import os
import sys
import openpyxl
from datetime import datetime

# Add root directory to python path
sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from backend.app.core.database import engine, Base, SessionLocal
from backend.app.models.models import User, Standard, Document, DocumentChunk, Laboratory
from backend.app.core.security import get_password_hash
from backend.app.rag.knowledge_seed import SEED_DOCUMENTS, SEED_LABORATORIES

CATEGORY_KEYWORDS = {
    "Cement & Concrete": ["cement", "concrete", "mortar", "aggregates", "pozzolana", "clinker"],
    "Electrical & Electronics": ["electric", "cable", "motor", "switch", "lamp", "transformer", "inverter", "fan", "battery", "insulation"],
    "Food & Agriculture": ["food", "fertilizer", "seed", "crop", "thresher", "water", "tea", "grain", "sugar", "beverage"],
    "Chemicals & Petrochemicals": ["chemical", "acid", "powder", "gas", "petroleum", "polymer", "plastic", "rubber", "explosive", "nitrate"],
    "Mechanical & Metallurgy": ["steel", "iron", "pipe", "valve", "bearing", "welding", "fastener", "cylinder", "rebar", "boiler"],
    "Civil & Construction": ["building", "brick", "tile", "timber", "glass", "roofing", "soil", "plumbing", "sanitary"],
    "Textiles": ["textile", "cotton", "yarn", "fabric", "wool", "silk", "garment"],
    "Medical & Healthcare": ["medical", "surgical", "syringe", "glove", "implant", "thermometer", "mask", "pharmaceutical"],
    "Consumer Goods & Safety": ["helmet", "toy", "jewellery", "gold", "silver", "cooker", "shoes", "footwear"]
}

def guess_category(title: str) -> str:
    title_lower = title.lower()
    for cat, keywords in CATEGORY_KEYWORDS.items():
        if any(kw in title_lower for kw in keywords):
            return cat
    return "General Engineering & Technical"

def seed_database():
    print("=" * 60)
    print("  BIS SmartAssist — Database Initialization & Seeding")
    print("=" * 60)
    
    # 1. Create Tables
    # Set stdout encoding for Windows
    if sys.stdout.encoding != 'utf-8':
        try:
            sys.stdout.reconfigure(encoding='utf-8')
        except Exception:
            pass

    print("\n[1/5] Creating database tables...")
    Base.metadata.create_all(bind=engine)
    print("[OK] Tables created successfully.")

    db = SessionLocal()
    try:
        # 2. Seed Users
        print("\n[2/5] Seeding initial users...")
        admin = db.query(User).filter(User.email == "admin@bis.gov.in").first()
        if not admin:
            admin_user = User(
                email="admin@bis.gov.in",
                hashed_password=get_password_hash("admin123"),
                full_name="BIS Nodal Officer (Admin)",
                role="admin",
                organization="Bureau of Indian Standards"
            )
            db.add(admin_user)

        demo_user = db.query(User).filter(User.email == "user@bis.gov.in").first()
        if not demo_user:
            regular_user = User(
                email="user@bis.gov.in",
                hashed_password=get_password_hash("user123"),
                full_name="Industry Applicant Demo",
                role="user",
                organization="MSME Manufacturing Consortium"
            )
            db.add(regular_user)
        db.commit()
        print("[OK] Users seeded: admin@bis.gov.in (Admin), user@bis.gov.in (User)")

        # 3. Ingest Curated Clause-level Documents
        print("\n[3/5] Ingesting curated BIS clause documents...")
        for doc_data in SEED_DOCUMENTS:
            existing_doc = db.query(Document).filter(Document.document_id == doc_data["document_id"]).first()
            if not existing_doc:
                new_doc = Document(
                    document_id=doc_data["document_id"],
                    title=doc_data["title"],
                    document_type=doc_data["document_type"],
                    standard_number=doc_data["standard_number"],
                    version=doc_data["version"],
                    source_url=doc_data["source_url"],
                    status="Indexed",
                    total_chunks=len(doc_data["chunks"])
                )
                db.add(new_doc)
                db.commit()

                for chunk_data in doc_data["chunks"]:
                    chunk = DocumentChunk(
                        document_id=doc_data["document_id"],
                        standard_number=doc_data["standard_number"],
                        title=doc_data["title"],
                        clause_number=chunk_data.get("clause_number"),
                        sub_clause=chunk_data.get("sub_clause"),
                        page_number=chunk_data.get("page_number"),
                        product_category=chunk_data.get("product_category"),
                        content=chunk_data["content"]
                    )
                    db.add(chunk)
                db.commit()
        print(f"[OK] Ingested {len(SEED_DOCUMENTS)} deep clause-level BIS reference documents.")

        # 4. Ingest BIS Laboratories
        print("\n[4/5] Ingesting BIS-recognized testing laboratories...")
        for lab_data in SEED_LABORATORIES:
            existing_lab = db.query(Laboratory).filter(Laboratory.lab_name == lab_data["lab_name"]).first()
            if not existing_lab:
                lab = Laboratory(**lab_data)
                db.add(lab)
        db.commit()
        print(f"[OK] Ingested {len(SEED_LABORATORIES)} BIS Regional Testing Laboratories.")

        # 5. Ingest Published Indian Standards from Excel
        print("\n[5/5] Ingesting published Indian Standards from Excel catalogue...")
        standards_count = db.query(Standard).count()
        excel_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "File_Published_Standards_List_2026-09-29_122533.xlsx")
        
        if standards_count < 1000 and os.path.exists(excel_path):
            print(f"Reading {excel_path}...")
            wb = openpyxl.load_workbook(excel_path, read_only=True)
            sheet = wb.active
            rows_iter = sheet.iter_rows(values_only=True)
            
            # Skip header rows
            next(rows_iter, None) # Row 0
            next(rows_iter, None) # Row 1 (column names)

            batch = []
            total_loaded = 0
            
            for row in rows_iter:
                if not row or not row[1]:
                    continue
                
                std_num = str(row[1]).strip()
                pub_date = str(row[2]).strip() if row[2] else None
                raw_title = str(row[3]).strip() if row[3] else "Standard Specification"
                clean_title = raw_title.replace("\ufffd", " - ").strip()
                std_type = str(row[4]).strip() if row[4] else "Specification"
                equiv = str(row[5]).strip() if row[5] else "Indigenous"
                category = guess_category(clean_title)

                is_mandatory = any(kw in clean_title.lower() for kw in ["safety", "protective", "drinking water", "cement", "steel", "helmet", "electric"])

                batch.append({
                    "standard_number": std_num,
                    "title": clean_title,
                    "date_of_publish": pub_date,
                    "type_of_standard": std_type,
                    "degree_of_equivalence": equiv,
                    "product_category": category,
                    "is_mandatory": is_mandatory,
                    "certification_scheme": "Scheme I (ISI Mark)",
                    "summary": f"Official Indian Standard for {clean_title}. Published by Bureau of Indian Standards.",
                    "key_requirements": "Conformity to specified chemical, physical, and performance criteria.",
                    "source_url": "https://www.services.bis.gov.in/php/BIS_2.0/bisconnect/knowyourstandards/indian_standards/"
                })

                if len(batch) >= 2000:
                    db.bulk_insert_mappings(Standard, batch)
                    db.commit()
                    total_loaded += len(batch)
                    print(f"   -> Inserted {total_loaded} standards...")
                    batch = []

            if batch:
                db.bulk_insert_mappings(Standard, batch)
                db.commit()
                total_loaded += len(batch)

            print(f"[OK] Successfully indexed all {total_loaded} published Indian Standards into database!")
        else:
            print(f"[OK] Standards already loaded ({standards_count} records present).")

        print("\n" + "=" * 60)
        print("  Database Seeding Completed Successfully!")
        print("=" * 60)

    except Exception as e:
        db.rollback()
        print(f"Error during seeding: {e}")
        raise e
    finally:
        db.close()

if __name__ == "__main__":
    seed_database()
